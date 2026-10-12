from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import Any

from charset_normalizer import from_bytes
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".html", ".htm", ".log", ".srt", ".vtt"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx"}
MEDIA_EXTENSIONS = {".mp4", ".mkv", ".mov", ".webm", ".m4v", ".mp3", ".m4a", ".wav", ".flac", ".aac", ".ogg"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | MEDIA_EXTENSIONS
IGNORED_PARTS = {".git", ".svn", "node_modules", ".venv", "__pycache__", "desktop_dist", "installer_dist"}
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_TOTAL_BYTES = 2 * 1024 * 1024 * 1024
MAX_FILES = 500
MAX_EXTRACTED_CHARS = 1_000_000
LARGE_TEXT_SAMPLE_BYTES = 1_000_000


@dataclass
class ExtractedDocument:
    text: str = ""
    status: str = "ready"
    error: str | None = None
    page_count: int | None = None
    slide_count: int | None = None
    sheet_count: int | None = None
    duration_seconds: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


class IgnoredRelativePathError(ValueError):
    """A valid relative path that is intentionally excluded from directory imports."""

    def __init__(self, relative_path: str):
        self.relative_path = relative_path
        super().__init__("已自动跳过隐藏目录或开发目录中的文件")


def safe_relative_path(value: str) -> str:
    value = (value or "unnamed").replace("\\", "/").strip("/")
    path = PurePosixPath(value)
    if not value or path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise ValueError("文件相对路径无效")
    relative_path = str(path)
    if any(part.lower() in IGNORED_PARTS or part.startswith(".") for part in path.parts[:-1]):
        raise IgnoredRelativePathError(relative_path)
    return relative_path


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    match = from_bytes(data).best()
    if not match:
        raise ValueError("无法识别文本编码")
    return str(match)


def strip_html(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    value = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h[1-6]>", "\n", value)
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    entities = {"&nbsp;": " ", "&amp;": "&", "&lt;": "<", "&gt;": ">", "&quot;": '"', "&#39;": "'"}
    for old, new in entities.items():
        value = value.replace(old, new)
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", value)).strip()


def ffprobe_path() -> str | None:
    candidates = []
    configured = os.getenv("FFPROBE_PATH")
    if configured:
        candidates.append(Path(configured))
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[3]))
    candidates.extend([root / "tools" / "ffmpeg" / "bin" / "ffprobe.exe", root / "tools" / "ffprobe.exe"])
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    return "ffprobe" if os.name != "nt" else None


def media_duration(data: bytes, suffix: str) -> int:
    executable = ffprobe_path()
    if not executable:
        raise RuntimeError("未找到 FFprobe，无法读取媒体时长")
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
            handle.write(data)
            temp_path = handle.name
        result = subprocess.run(
            [executable, "-v", "error", "-show_entries", "format=duration", "-of", "json", temp_path],
            capture_output=True, text=True, timeout=60, check=True,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        seconds = float(json.loads(result.stdout)["format"]["duration"])
        return max(1, round(seconds))
    finally:
        if temp_path:
            Path(temp_path).unlink(missing_ok=True)


def media_duration_path(path: Path) -> int:
    """Read media metadata without copying a potentially huge file."""
    executable = ffprobe_path()
    if not executable:
        raise RuntimeError("未找到 FFprobe，无法读取媒体时长")
    result = subprocess.run(
        [executable, "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)],
        capture_output=True, text=True, timeout=120, check=True,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
    )
    seconds = float(json.loads(result.stdout)["format"]["duration"])
    return max(1, round(seconds))


def sha256_path(path: Path, block_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(block_size):
            digest.update(chunk)
    return digest.hexdigest()


def sampled_text(path: Path, max_bytes: int = LARGE_TEXT_SAMPLE_BYTES) -> tuple[str, bool]:
    """Return bounded, representative text from across a large file.

    Reading evenly spaced windows avoids loading multi-gigabyte learning material
    into RAM while still exposing its beginning, middle sections and ending to the
    planner. Byte offsets are included so the sample remains auditable.
    """
    size = path.stat().st_size
    if size <= max_bytes:
        data = path.read_bytes()
        return decode_text(data), False
    window_count = 12
    window_size = max(16 * 1024, max_bytes // window_count)
    with path.open("rb") as handle:
        handle.seek(0)
        head = handle.read(min(256 * 1024, window_size))
        match = from_bytes(head).best()
        encoding = match.encoding if match and match.encoding else "utf-8"
        pieces = []
        max_offset = max(0, size - window_size)
        for index in range(window_count):
            offset = round(max_offset * index / max(1, window_count - 1))
            handle.seek(offset)
            data = handle.read(window_size)
            text = data.decode(encoding, errors="replace")
            if offset:
                newline = text.find("\n")
                if 0 <= newline < 4096:
                    text = text[newline + 1:]
            pieces.append(f"[文件字节位置 {offset:,} / {size:,}]\n{text.strip()}")
    return "\n\n".join(pieces)[:MAX_EXTRACTED_CHARS], True


def extract_document_path(filename: str, path: Path) -> ExtractedDocument:
    """Extract from a filesystem path with bounded memory usage."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return ExtractedDocument(status="skipped", error=f"暂不支持 {suffix or '无扩展名'} 文件")
    try:
        if suffix in TEXT_EXTENSIONS:
            text, sampled = sampled_text(path)
            if suffix in {".html", ".htm"}:
                text = strip_html(text)
            result = ExtractedDocument(text=text, metadata={"text_sampled": sampled})
        elif suffix == ".pdf":
            reader = PdfReader(str(path))
            pages = [f"[第 {index + 1} 页]\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages)]
            result = ExtractedDocument(text="\n\n".join(pages), page_count=len(reader.pages))
        elif suffix == ".docx":
            doc = Document(str(path))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            result = ExtractedDocument(text="\n".join(parts))
        elif suffix == ".pptx":
            deck = Presentation(str(path))
            slides = []
            for index, slide in enumerate(deck.slides):
                lines = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                slides.append(f"[第 {index + 1} 页]\n" + "\n".join(lines))
            result = ExtractedDocument(text="\n\n".join(slides), slide_count=len(deck.slides))
        elif suffix == ".xlsx":
            book = load_workbook(str(path), read_only=True, data_only=True)
            sheets = []
            for sheet in book.worksheets:
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        lines.append(" | ".join(values))
                sheets.append(f"[工作表：{sheet.title}]\n" + "\n".join(lines))
            result = ExtractedDocument(text="\n\n".join(sheets), sheet_count=len(book.worksheets))
        else:
            duration = media_duration_path(path)
            result = ExtractedDocument(duration_seconds=duration, metadata={"media_type": "audio" if suffix in {".mp3", ".m4a", ".wav", ".flac", ".aac", ".ogg"} else "video"})
        if len(result.text) > MAX_EXTRACTED_CHARS:
            result.text = result.text[:MAX_EXTRACTED_CHARS]
            result.metadata["text_truncated"] = True
        return result
    except Exception as exc:
        return ExtractedDocument(status="error", error=f"{exc.__class__.__name__}: {exc}")


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return ExtractedDocument(status="skipped", error=f"暂不支持 {suffix or '无扩展名'} 文件")
    try:
        if suffix in {".txt", ".md", ".markdown", ".log", ".srt", ".vtt"}:
            result = ExtractedDocument(text=decode_text(data))
        elif suffix == ".json":
            payload = json.loads(decode_text(data))
            result = ExtractedDocument(text=json.dumps(payload, ensure_ascii=False, indent=2))
        elif suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            rows = csv.reader(io.StringIO(decode_text(data)), delimiter=delimiter)
            result = ExtractedDocument(text="\n".join(" | ".join(cell.strip() for cell in row) for row in rows))
        elif suffix in {".html", ".htm"}:
            result = ExtractedDocument(text=strip_html(decode_text(data)))
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            pages = [f"[第 {index + 1} 页]\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages)]
            result = ExtractedDocument(text="\n\n".join(pages), page_count=len(reader.pages))
            if not result.text.replace("[", "").strip():
                result.status, result.error = "warning", "PDF 可能是扫描件，未提取到正文"
        elif suffix == ".docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            result = ExtractedDocument(text="\n".join(parts))
        elif suffix == ".pptx":
            deck = Presentation(io.BytesIO(data))
            slides = []
            for index, slide in enumerate(deck.slides):
                lines = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                slides.append(f"[第 {index + 1} 页]\n" + "\n".join(lines))
            result = ExtractedDocument(text="\n\n".join(slides), slide_count=len(deck.slides))
        elif suffix == ".xlsx":
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = []
            for sheet in book.worksheets:
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        lines.append(" | ".join(values))
                sheets.append(f"[工作表：{sheet.title}]\n" + "\n".join(lines))
            result = ExtractedDocument(text="\n\n".join(sheets), sheet_count=len(book.worksheets))
        else:
            duration = media_duration(data, suffix)
            result = ExtractedDocument(duration_seconds=duration, metadata={"media_type": "audio" if suffix in {".mp3", ".m4a", ".wav", ".flac", ".aac", ".ogg"} else "video"})
        if len(result.text) > MAX_EXTRACTED_CHARS:
            result.text = result.text[:MAX_EXTRACTED_CHARS]
            result.metadata["text_truncated"] = True
        return result
    except Exception as exc:
        return ExtractedDocument(status="error", error=f"{exc.__class__.__name__}: {exc}")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
