import io
import json

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.schemas import LearningUnit
from app.services.planning import _partition_ordered


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("第二章：概率基础", level=1)
    document.add_paragraph("条件概率、贝叶斯公式与三个练习题。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pptx_bytes() -> bytes:
    deck = Presentation()
    for title in ("模型评估", "交叉验证", "误差分析"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"{title}的关键步骤和示例"
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "实验安排"
    sheet.append(["实验", "预计分钟", "前置内容"])
    sheet.append(["数据清洗", 45, "概率基础"])
    sheet.append(["训练模型", 70, "数据清洗"])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_ordered_partition_keeps_exact_groups_and_minimizes_peak_load():
    def make_units(weights):
        return [LearningUnit(
            id=f"u-{index}", title=f"单元 {index}", document_id=index,
            source_path=f"{index}.md", sequence=index, estimated_minutes=weight,
        ) for index, weight in enumerate(weights)]

    skewed = _partition_ordered(make_units([100, 1, 1]), 3)
    assert [[unit.estimated_minutes for unit in day] for day in skewed] == [[100], [1], [1]]

    weights = [45, 80, 75, 57, 56, 53, 52, 55, 20]
    balanced = _partition_ordered(make_units(weights), 7)
    flattened = [unit.estimated_minutes for day in balanced for unit in day]
    assert flattened == weights
    assert all(balanced)
    assert max(sum(unit.estimated_minutes for unit in day) for day in balanced) == 105


def test_complex_multiformat_course_is_sanitized_balanced_and_confirmed(client, monkeypatch):
    """A realistic course catches malformed LLM ranges, ordering and review rounding."""
    monkeypatch.setenv("COMPLEX_PLANNER_KEY", "sk-complex-test")
    monkeypatch.setattr("app.services.document_import.media_duration", lambda _data, _suffix: 5_400)
    client.post("/api/settings/providers", json={
        "name": "Complex Planner",
        "base_url": "https://example.invalid/v1",
        "model": "planner-complex-model",
        "api_key_env": "COMPLEX_PLANNER_KEY",
        "enabled": True,
    })

    paths = [
        "机器学习课程/01-课程导论.md",
        "机器学习课程/02-概率基础.docx",
        "机器学习课程/03-模型评估.pptx",
        "机器学习课程/04-实验安排.xlsx",
        "机器学习课程/05-完整演示.mp4",
        "机器学习课程/06-实验讲解.srt",
        "机器学习课程/07-扩展阅读.html",
    ]
    response = client.post(
        "/api/import/batches",
        data={"root_name": "机器学习课程", "source_kind": "directory", "relative_paths": json.dumps(paths, ensure_ascii=False)},
        files=[
            ("files", ("01-课程导论.md", "# 课程导论\n先理解学习路线、符号和最终项目。".encode(), "text/markdown")),
            ("files", ("02-概率基础.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("03-模型评估.pptx", _pptx_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
            ("files", ("04-实验安排.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ("files", ("05-完整演示.mp4", b"fake-media-for-mocked-ffprobe", "video/mp4")),
            ("files", ("06-实验讲解.srt", b"1\n00:00:00,000 --> 00:01:00,000\nPrepare the dataset.\n", "application/x-subrip")),
            ("files", ("07-扩展阅读.html", "<h1>扩展阅读</h1><p>比较不同评估指标。</p>".encode(), "text/html")),
        ],
    )
    assert response.status_code == 200, response.text
    batch = response.json()
    assert batch["supported_count"] == 7
    assert batch["error_count"] == 0
    assert batch["total_media_seconds"] == 5_400
    ids = {doc["filename"]: doc["id"] for doc in batch["documents"]}

    captured_request = {}

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            # Deliberately unordered, with duplicate IDs and impossible source ranges.
            units = [
                {"id": "duplicate", "title": "完成训练实验", "document_id": ids["04-实验安排.xlsx"], "source_path": "wrong.xlsx", "unit_type": "project", "topic": "实验", "sequence": 4, "estimated_minutes": 105},
                {"id": "intro", "title": "理解课程路线", "document_id": ids["01-课程导论.md"], "source_path": "wrong.md", "unit_type": "reading", "topic": "导论", "sequence": 0, "estimated_minutes": 45},
                {"id": "duplicate", "title": "学习概率基础", "document_id": ids["02-概率基础.docx"], "source_path": "wrong.docx", "unit_type": "reading", "topic": "概率", "sequence": 1, "estimated_minutes": 80},
                {"id": "metrics", "title": "掌握模型评估", "document_id": ids["03-模型评估.pptx"], "source_path": "wrong.pptx", "unit_type": "reading", "topic": "评估", "sequence": 2, "estimated_minutes": 75, "page_from": -4, "page_to": 99},
                {"id": "video", "title": "观看完整演示", "document_id": ids["05-完整演示.mp4"], "source_path": "wrong.mp4", "unit_type": "video", "topic": "演示", "sequence": 3, "estimated_minutes": 999, "media_start": -20, "media_end": 99_999},
                {"id": "lab", "title": "跟做字幕实验", "document_id": ids["06-实验讲解.srt"], "source_path": "wrong.srt", "unit_type": "exercise", "topic": "实验", "sequence": 5, "estimated_minutes": 55},
                # The HTML document is intentionally omitted to exercise local coverage repair.
            ]
            return {"choices": [{"message": {"content": json.dumps({
                "title": "机器学习七日进阶计划",
                "goal": "完成理论、演示和实验闭环",
                "units": units,
            }, ensure_ascii=False)}}]}

    def fake_post(*args, **kwargs):
        captured_request.update(kwargs.get("json", {}))
        return FakeResponse()

    monkeypatch.setattr("app.services.planning.httpx.post", fake_post)
    scheduled = client.post("/api/planning/schedule", json={
        "batch_id": batch["id"],
        "start_date": "2026-10-05",
        "day_count": 7,
        "include_weekends": False,
        "excluded_dates": ["2026-10-08"],
        "daily_available_minutes": 120,
        "intensity": "standard",
        "review_ratio": 0.15,
        "buffer_ratio": 0.1,
        "goal": "在七个学习日完成课程并产出实验结果",
    })
    assert scheduled.status_code == 200, scheduled.text
    schedule = scheduled.json()

    assert [day["date"] for day in schedule["days"]] == [
        "2026-10-05", "2026-10-06", "2026-10-07", "2026-10-09",
        "2026-10-12", "2026-10-13", "2026-10-14",
    ]
    assert schedule["provider"] == "Complex Planner"
    assert len(captured_request["messages"]) == 2
    prompt = captured_request["messages"][1]["content"]
    assert all(path in prompt for path in paths)

    all_units = [unit for day in schedule["days"] for unit in day["units"]]
    content_units = [unit for unit in all_units if unit["unit_type"] != "review"]
    review_units = [unit for unit in all_units if unit["unit_type"] == "review"]
    assert {unit["document_id"] for unit in content_units} == set(ids.values())
    assert len({unit["id"] for unit in all_units}) == len(all_units)
    assert any("补回模型遗漏" in warning for warning in schedule["warnings"])

    # The planner, not the LLM, is authoritative for real media duration and ranges.
    video_parts = [unit for unit in content_units if unit["document_id"] == ids["05-完整演示.mp4"]]
    assert video_parts[0]["media_start"] == 0
    assert video_parts[-1]["media_end"] == 5_400
    assert all(left["media_end"] == right["media_start"] for left, right in zip(video_parts, video_parts[1:]))
    assert sum(unit["estimated_minutes"] for unit in video_parts) == 113  # ceil(90 min * standard 1.25)

    slide_units = [unit for unit in content_units if unit["document_id"] == ids["03-模型评估.pptx"]]
    assert min(unit["page_from"] for unit in slide_units) >= 1
    assert max(unit["page_to"] for unit in slide_units) <= 3

    base_minutes = sum(unit["estimated_minutes"] for unit in content_units)
    expected_review = round(base_minutes * 0.15)
    assert sum(unit["estimated_minutes"] for unit in review_units) == expected_review
    assert schedule["total_estimated_minutes"] == base_minutes + expected_review
    assert all(day["units"] for day in schedule["days"])
    assert all(day["estimated_minutes"] <= day["capacity_minutes"] for day in schedule["days"]), [
        (day["date"], day["estimated_minutes"], day["capacity_minutes"]) for day in schedule["days"]
    ]

    confirmed = client.post("/api/planning/confirm", json={"batch_id": batch["id"], "schedule": schedule})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["created"] == len(schedule["days"]) + len(all_units)
    saved = client.get("/api/items").json()
    parents = [item for item in saved if item["parent_id"] is None and item["source"] == "Directory:机器学习课程"]
    assert len(parents) == 7
    assert [item["due_date"] for item in sorted(parents, key=lambda item: item["due_date"])] == [day["date"] for day in schedule["days"]]
