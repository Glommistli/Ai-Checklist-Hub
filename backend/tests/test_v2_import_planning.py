import io
import json

from app.schemas import LearningUnit, LearningUnitSet, PlanningOptions
from app.services.document_import import extract_document, safe_relative_path, sampled_text
from app.services.planning import build_schedule


def test_text_json_and_office_extractors():
    assert "中文" in extract_document("notes.txt", "中文资料".encode()).text
    assert '"title": "plan"' in extract_document("data.json", b'{"title":"plan"}').text

    from docx import Document
    doc = Document(); doc.add_heading("第一章", level=1); doc.add_paragraph("核心概念")
    stream = io.BytesIO(); doc.save(stream)
    assert "核心概念" in extract_document("course.docx", stream.getvalue()).text

    from pptx import Presentation
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = "课程导论"
    stream = io.BytesIO(); deck.save(stream)
    ppt = extract_document("course.pptx", stream.getvalue())
    assert ppt.slide_count == 1 and "课程导论" in ppt.text

    from openpyxl import Workbook
    book = Workbook(); book.active["A1"] = "章节"; book.active["B1"] = "时长"
    stream = io.BytesIO(); book.save(stream)
    sheet = extract_document("plan.xlsx", stream.getvalue())
    assert sheet.sheet_count == 1 and "章节 | 时长" in sheet.text


def test_safe_relative_path_blocks_traversal_and_dev_directories():
    assert safe_relative_path("课程/第一章.md") == "课程/第一章.md"
    for unsafe in ("../secret.txt", ".git/config", "course/node_modules/a.js"):
        try:
            safe_relative_path(unsafe)
            assert False, unsafe
        except ValueError:
            pass


def test_batch_upload_extracts_multiple_files(client):
    response = client.post(
        "/api/import/batches",
        data={"root_name": "学习资料", "source_kind": "directory", "relative_paths": json.dumps(["学习资料/01.md", "学习资料/02.json"])},
        files=[("files", ("01.md", b"# First\nRead chapter one", "text/markdown")), ("files", ("02.json", b'{"task":"practice"}', "application/json"))],
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["file_count"] == 2
    assert payload["supported_count"] == 2
    assert [doc["relative_path"] for doc in payload["documents"]] == ["学习资料/01.md", "学习资料/02.json"]


def test_large_manifest_avoids_copying_47gb_media(client, monkeypatch, tmp_path):
    monkeypatch.setattr("app.api.routes._upload_root", lambda: tmp_path)
    response = client.post("/api/import/uploads", json={
        "root_name": "大型课程", "source_kind": "directory", "files": [
            {"relative_path": "大型课程/三GB讲义.txt", "size": 3 * 1024**3, "mime_type": "text/plain"},
            {"relative_path": "大型课程/四十七GB视频.mp4", "size": 47 * 1024**3, "mime_type": "video/mp4", "duration_seconds": 8 * 3600},
        ],
    })
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["total_size"] == 50 * 1024**3
    text_doc, video_doc = payload["documents"]
    assert text_doc["extraction_status"] == "pending"
    assert text_doc["metadata"]["upload_required"] is True
    assert video_doc["extraction_status"] == "ready"
    assert video_doc["duration_seconds"] == 8 * 3600
    assert video_doc["metadata"]["metadata_only"] is True
    assert not (tmp_path / str(payload["id"])).exists()


def test_chunked_directory_skips_hidden_and_dev_paths_without_aborting(client, monkeypatch, tmp_path):
    monkeypatch.setattr("app.api.routes._upload_root", lambda: tmp_path)
    content = b"# Lesson\nReview the chapter"
    response = client.post("/api/import/uploads", json={
        "root_name": "课程资料",
        "source_kind": "directory",
        "files": [
            {"relative_path": "课程资料/讲义.md", "size": len(content), "mime_type": "text/markdown"},
            {"relative_path": "课程资料/.git/config", "size": 200, "mime_type": "text/plain"},
            {"relative_path": "课程资料/node_modules/readme.md", "size": 300, "mime_type": "text/markdown"},
        ],
    })
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["file_count"] == 3
    assert payload["skipped_count"] == 2
    assert len(payload["documents"]) == 3
    assert payload["documents"][0]["extraction_status"] == "pending"
    for document in payload["documents"][1:]:
        assert document["extraction_status"] == "skipped"
        assert document["metadata"]["upload_required"] is False
        assert document["metadata"]["ignored_path"] is True
        assert "已自动跳过" in document["error_message"]

    batch_id = payload["id"]
    ready_document = payload["documents"][0]
    uploaded = client.put(
        f"/api/import/uploads/{batch_id}/{ready_document['id']}/chunk?offset=0",
        content=content,
    )
    assert uploaded.status_code == 200, uploaded.text
    extracted = client.post(f"/api/import/uploads/{batch_id}/{ready_document['id']}/complete")
    assert extracted.status_code == 200, extracted.text
    completed = client.post(f"/api/import/uploads/{batch_id}/complete")
    assert completed.status_code == 200, completed.text
    final_payload = completed.json()
    assert final_payload["status"] == "ready"
    assert final_payload["supported_count"] == 1
    assert final_payload["skipped_count"] == 2
    assert final_payload["error_count"] == 0


def test_chunked_directory_still_rejects_path_traversal(client, monkeypatch, tmp_path):
    monkeypatch.setattr("app.api.routes._upload_root", lambda: tmp_path)
    response = client.post("/api/import/uploads", json={
        "root_name": "课程资料",
        "source_kind": "directory",
        "files": [{"relative_path": "../secret.md", "size": 10, "mime_type": "text/markdown"}],
    })
    assert response.status_code == 400
    assert "路径无效" in response.json()["detail"]


def test_chunked_upload_retries_and_completes(client, monkeypatch, tmp_path):
    monkeypatch.setattr("app.api.routes._upload_root", lambda: tmp_path)
    content = "第一章：概念\n第二章：练习".encode()
    started = client.post("/api/import/uploads", json={
        "root_name": "课程", "source_kind": "files",
        "files": [{"relative_path": "课程.md", "size": len(content), "mime_type": "text/markdown"}],
    })
    assert started.status_code == 201, started.text
    batch = started.json(); document = batch["documents"][0]
    split = len(content) // 2
    first = client.put(f"/api/import/uploads/{batch['id']}/{document['id']}/chunk?offset=0", content=content[:split])
    assert first.status_code == 200, first.text
    replay = client.put(f"/api/import/uploads/{batch['id']}/{document['id']}/chunk?offset=0", content=content[:split])
    assert replay.status_code == 200 and replay.json()["replayed"] is True
    second = client.put(f"/api/import/uploads/{batch['id']}/{document['id']}/chunk?offset={split}", content=content[split:])
    assert second.status_code == 200, second.text
    extracted = client.post(f"/api/import/uploads/{batch['id']}/{document['id']}/complete")
    assert extracted.status_code == 200, extracted.text
    completed = client.post(f"/api/import/uploads/{batch['id']}/complete")
    assert completed.status_code == 200, completed.text
    payload = completed.json()
    assert payload["status"] == "ready"
    assert payload["supported_count"] == 1
    assert not (tmp_path / str(batch["id"])).exists()


def test_three_gb_sparse_text_is_sampled_with_bounded_read(tmp_path):
    path = tmp_path / "three-gb-notes.txt"
    size = 3 * 1024**3
    with path.open("wb") as handle:
        handle.write("开头章节\n".encode())
        handle.seek(size - 32)
        handle.write("结尾复习\n".encode())
    text, was_sampled = sampled_text(path, max_bytes=256 * 1024)
    assert path.stat().st_size > 3_000_000_000
    assert was_sampled is True
    assert len(text) <= 1_000_000
    assert "文件字节位置" in text


def test_planning_requires_enabled_llm(client):
    batch = client.post(
        "/api/import/batches", data={"root_name": "资料", "relative_paths": '["a.md"]'},
        files=[("files", ("a.md", b"# Topic\nStudy it", "text/markdown"))],
    ).json()
    response = client.post("/api/planning/schedule", json={
        "batch_id": batch["id"], "start_date": "2026-09-23", "day_count": 3,
        "daily_available_minutes": 60, "goal": "完成学习",
    })
    assert response.status_code == 400
    assert "强制使用 LLM" in response.json()["detail"]


def test_llm_schedule_and_confirm_flow(client, monkeypatch):
    monkeypatch.setenv("PLANNER_TEST_KEY", "sk-test")
    client.post("/api/settings/providers", json={
        "name": "Planner", "base_url": "https://example.invalid/v1", "model": "planner-model",
        "api_key_env": "PLANNER_TEST_KEY", "enabled": True,
    })
    batch = client.post(
        "/api/import/batches", data={"root_name": "课程", "relative_paths": '["01.md","02.md"]'},
        files=[("files", ("01.md", b"chapter one", "text/markdown")), ("files", ("02.md", b"chapter two", "text/markdown"))],
    ).json()

    class FakeResponse:
        def raise_for_status(self): pass
        def json(self):
            return {"choices": [{"message": {"content": json.dumps({
                "title": "三天课程计划", "goal": "掌握两章内容", "units": [
                    {"id": "u1", "title": "学习第一章", "document_id": batch["documents"][0]["id"], "source_path": "01.md", "unit_type": "reading", "topic": "基础", "sequence": 0, "estimated_minutes": 80},
                    {"id": "u2", "title": "学习第二章", "document_id": batch["documents"][1]["id"], "source_path": "02.md", "unit_type": "exercise", "topic": "进阶", "sequence": 1, "estimated_minutes": 100},
                ]
            }, ensure_ascii=False)}}]}

    monkeypatch.setattr("app.services.planning.httpx.post", lambda *args, **kwargs: FakeResponse())
    scheduled = client.post("/api/planning/schedule", json={
        "batch_id": batch["id"], "start_date": "2026-09-23", "day_count": 3,
        "daily_available_minutes": 90, "review_ratio": 0,
    })
    assert scheduled.status_code == 200, scheduled.text
    schedule = scheduled.json()
    assert schedule["goal"] == "掌握两章内容"
    assert len(schedule["days"]) == 3
    assert schedule["provider"] == "Planner"
    assigned = [unit["id"] for day in schedule["days"] for unit in day["units"]]
    assert assigned == ["u1", "u2-part-1", "u2-part-2"]

    confirmed = client.post("/api/planning/confirm", json={"batch_id": batch["id"], "schedule": schedule})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["created"] == 6
    saved = client.get("/api/items").json()
    parents = [item for item in saved if item["parent_id"] is None and item["source"] == "Directory:课程"]
    assert len(parents) == 3
    assert {item["due_date"] for item in parents} == {"2026-09-23", "2026-09-24", "2026-09-25"}


def test_scheduler_skips_weekends_and_excluded_dates():
    units = LearningUnitSet(title="计划", units=[LearningUnit(
        id="u", title="学习", document_id=1, source_path="a.md", estimated_minutes=60,
    )])
    preview = build_schedule(units, PlanningOptions(
        batch_id=1, start_date="2026-09-25", day_count=3, include_weekends=False,
        excluded_dates=["2026-09-29"], daily_available_minutes=60, goal="学习",
    ), "Provider", "Model")
    assert [str(day.date) for day in preview.days] == ["2026-09-25", "2026-09-28", "2026-09-30"]


def test_scheduler_supports_start_and_end_date_mode():
    units = LearningUnitSet(title="日期范围计划", units=[LearningUnit(
        id="u", title="学习", document_id=1, source_path="a.md", estimated_minutes=60,
    )])
    preview = build_schedule(units, PlanningOptions(
        batch_id=1, start_date="2026-09-25", schedule_mode="end_date", end_date="2026-09-30",
        include_weekends=False, excluded_dates=["2026-09-29"], daily_available_minutes=60,
    ), "Provider", "Model")
    assert [str(day.date) for day in preview.days] == ["2026-09-25", "2026-09-28", "2026-09-30"]
    assert str(preview.start_date) == "2026-09-25"
    assert str(preview.end_date) == "2026-09-30"
    assert preview.day_count == 3


def test_over_capacity_schedule_preserves_every_unit():
    source_units = [LearningUnit(
        id=f"unit-{index}", title=f"学习单元 {index + 1}", document_id=1,
        source_path="large-course.md", topic=f"主题 {index // 4 + 1}", sequence=index,
        estimated_minutes=120, splittable=False,
    ) for index in range(20)]
    preview = build_schedule(LearningUnitSet(title="大型资料", units=source_units), PlanningOptions(
        batch_id=1, start_date="2026-10-01", day_count=3,
        daily_available_minutes=60, review_ratio=0,
    ), "Provider", "Model")
    assigned = [unit.id for day in preview.days for unit in day.units]
    assert assigned == [unit.id for unit in source_units]
    assert preview.unscheduled_units == []
    assert any("所有资料仍会保留并创建为可编辑任务" in warning for warning in preview.warnings)
